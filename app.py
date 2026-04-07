"""
Destek Home | Kargo Desi Otomasyon Sistemi
Production-ready Streamlit app for daily e-commerce cargo operations.
UI Language: Turkish | Logic: Python/Pandas
"""

import io
import re
import streamlit as st
import pandas as pd
from datetime import datetime

# ─────────────────────────────────────────────
# PAGE CONFIG
# ─────────────────────────────────────────────
st.set_page_config(
    page_title="Kargo Desi Otomasyonu",
    page_icon="📦",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ─────────────────────────────────────────────
# CUSTOM CSS  — industrial/utilitarian palette
# ─────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:wght@400;600&family=IBM+Plex+Sans:wght@300;400;600;700&display=swap');

/* ── Root tokens ── */
:root {
    --bg:        #0f1117;
    --surface:   #1a1d27;
    --surface2:  #22263a;
    --border:    #2e3350;
    --accent:    #1A5C6B;   /* Okyanus Teal  */
    --accent2:   #D4652B;   /* Ember Turuncu */
    --cream:     #F0EDE8;   /* Sıcak Krem    */
    --ok:        #22c55e;
    --warn:      #f59e0b;
    --err:       #ef4444;
    --text:      #e2e8f0;
    --muted:     #64748b;
    --mono:      'IBM Plex Mono', monospace;
    --sans:      'IBM Plex Sans', sans-serif;
}

/* ── Base ── */
html, body, [class*="css"] { font-family: var(--sans); }
.stApp { background: var(--bg); color: var(--text); }

/* ── Header banner ── */
.hero-banner {
    background: linear-gradient(135deg, var(--accent) 0%, #0d3d4a 60%, #1a1d27 100%);
    border: 1px solid var(--border);
    border-radius: 12px;
    padding: 28px 36px;
    margin-bottom: 28px;
    position: relative;
    overflow: hidden;
}
.hero-banner::before {
    content: '📦';
    position: absolute;
    right: 36px; top: 50%;
    transform: translateY(-50%);
    font-size: 72px;
    opacity: 0.12;
}
.hero-banner h1 {
    font-size: 1.9rem;
    font-weight: 700;
    color: var(--cream);
    margin: 0 0 4px 0;
    letter-spacing: -0.5px;
}
.hero-banner p {
    color: #93c5d5;
    margin: 0;
    font-size: 0.92rem;
    font-weight: 300;
}
.hero-badge {
    display: inline-block;
    background: var(--accent2);
    color: white;
    font-size: 0.7rem;
    font-weight: 600;
    padding: 2px 10px;
    border-radius: 20px;
    letter-spacing: 1px;
    text-transform: uppercase;
    margin-bottom: 10px;
}

/* ── Upload cards ── */
.upload-card {
    background: var(--surface);
    border: 1px solid var(--border);
    border-radius: 10px;
    padding: 20px 24px 16px;
    margin-bottom: 4px;
}
.upload-card h4 {
    color: var(--cream);
    font-size: 0.95rem;
    font-weight: 600;
    margin: 0 0 4px 0;
}
.upload-card p {
    color: var(--muted);
    font-size: 0.8rem;
    margin: 0 0 12px 0;
    font-family: var(--mono);
}

/* ── Metric cards ── */
[data-testid="metric-container"] {
    background: var(--surface);
    border: 1px solid var(--border);
    border-radius: 10px;
    padding: 16px 20px;
}
[data-testid="stMetricValue"] {
    font-family: var(--mono);
    font-size: 2.2rem !important;
    color: var(--cream) !important;
}
[data-testid="stMetricLabel"] { color: var(--muted) !important; font-size: 0.78rem !important; }
[data-testid="stMetricDelta"] { font-family: var(--mono); font-size: 0.82rem !important; }

/* ── Tabs ── */
[data-testid="stTabs"] [data-baseweb="tab-list"] {
    background: var(--surface);
    border-radius: 8px;
    padding: 4px;
    gap: 4px;
    border: 1px solid var(--border);
}
[data-testid="stTabs"] [data-baseweb="tab"] {
    border-radius: 6px;
    color: var(--muted);
    font-weight: 600;
    font-size: 0.85rem;
}
[data-testid="stTabs"] [aria-selected="true"] {
    background: var(--accent) !important;
    color: white !important;
}

/* ── Dataframe ── */
[data-testid="stDataFrame"] { border-radius: 8px; overflow: hidden; }

/* ── Buttons ── */
.stDownloadButton > button {
    background: linear-gradient(135deg, var(--accent2), #b8541f) !important;
    color: white !important;
    border: none !important;
    border-radius: 8px !important;
    font-weight: 700 !important;
    font-size: 1rem !important;
    padding: 14px 32px !important;
    width: 100% !important;
    letter-spacing: 0.3px;
    transition: opacity 0.2s;
}
.stDownloadButton > button:hover { opacity: 0.88 !important; }

.stButton > button {
    background: var(--accent) !important;
    color: white !important;
    border: none !important;
    border-radius: 8px !important;
    font-weight: 600 !important;
}

/* ── Alert overrides ── */
[data-testid="stAlert"] { border-radius: 8px; }

/* ── Divider ── */
hr { border-color: var(--border) !important; margin: 24px 0; }

/* ── Status pill ── */
.pill-ok   { background:#14532d22; color:var(--ok);  border:1px solid var(--ok);  border-radius:20px; padding:2px 10px; font-size:0.75rem; font-family:var(--mono); }
.pill-warn { background:#78350f22; color:var(--warn); border:1px solid var(--warn); border-radius:20px; padding:2px 10px; font-size:0.75rem; font-family:var(--mono); }

/* ── Sidebar ── */
[data-testid="stSidebar"] { background: var(--surface); border-right: 1px solid var(--border); }
</style>
""", unsafe_allow_html=True)


# ─────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────

def clean_text(series: pd.Series) -> pd.Series:
    """Strip whitespace and lowercase."""
    return series.astype(str).str.strip().str.lower()


def validate_columns(df: pd.DataFrame, required: list[str], file_label: str) -> list[str]:
    """Return list of missing columns."""
    missing = [col for col in required if col not in df.columns]
    return missing


def find_desi(product_name: str, keyword_map: dict[str, float]) -> float | str:
    """
    Case-insensitive substring match.
    Longer keywords are checked first to avoid partial shadowing.
    Returns the Desi value or 'KONTROL ET'.
    """
    name_lower = product_name.lower().strip()
    # Sort keywords by length descending (most specific first)
    for kw in sorted(keyword_map.keys(), key=len, reverse=True):
        if kw in name_lower:
            return keyword_map[kw]
    return "KONTROL ET"


def build_excel(df: pd.DataFrame) -> bytes:
    """Return an in-memory Excel file as bytes."""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Kargo Listesi")
        ws = writer.sheets["Kargo Listesi"]

        # Auto-fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col) + 4
            ws.column_dimensions[col[0].column_letter].width = min(max_len, 60)

        # Style header row
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        header_fill = PatternFill("solid", fgColor="1A5C6B")
        header_font = Font(bold=True, color="F0EDE8", name="Calibri", size=11)
        thin = Side(style="thin", color="2e3350")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        # Highlight KONTROL ET rows in amber
        warn_fill = PatternFill("solid", fgColor="FEF3C7")
        warn_font = Font(bold=True, color="92400E", name="Calibri", size=10)
        normal_font = Font(name="Calibri", size=10)
        normal_fill = PatternFill("solid", fgColor="FFFFFF")

        for row in ws.iter_rows(min_row=2):
            desi_cell = row[1]  # Column B = Desi
            is_warn = str(desi_cell.value) == "KONTROL ET"
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if is_warn:
                    cell.fill = warn_fill
                    cell.font = warn_font
                else:
                    cell.fill = normal_fill
                    cell.font = normal_font

        # Freeze top row
        ws.freeze_panes = "A2"

    return output.getvalue()


# ─────────────────────────────────────────────
# SIDEBAR — instructions
# ─────────────────────────────────────────────
with st.sidebar:
    st.markdown("### 📋 Kullanım Kılavuzu")
    st.markdown("""
**Adım 1 — Desi Kılavuzu**  
Sütun A: `Anahtar Kelime`  
Sütun B: `Desi` (sayısal)

**Adım 2 — Günlük Siparişler**  
`Ürün Adı` ve `Kargo Kodu` sütunlarını içermeli.

**Adım 3 — İşle**  
"İşlemi Başlat" butonuna bas.

**Adım 4 — İndir**  
Hazırlanan kargo dosyasını Excel olarak indir.

---
**Eşleştirme Mantığı**  
Anahtar kelime, ürün adının herhangi bir yerinde aranır. Birden fazla eşleşme varsa en uzun (en spesifik) anahtar kelime kazanır.

---
<span style='color:#64748b; font-size:0.75rem;'>Destek Home · Kargo Otomasyon v2.0</span>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────────
# HERO BANNER
# ─────────────────────────────────────────────
st.markdown("""
<div class="hero-banner">
  <div class="hero-badge">Kargo Otomasyonu</div>
  <h1>📦 Desi Eşleştirme Sistemi</h1>
  <p>Günlük siparişlerinizi yükleyin · otomatik eşleştirin · kargo dosyasını indirin</p>
</div>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────────
# FILE UPLOADS
# ─────────────────────────────────────────────
col_up1, col_up2 = st.columns(2, gap="medium")

with col_up1:
    st.markdown("""
    <div class="upload-card">
      <h4>📗 Desi Kılavuzu</h4>
      <p>Anahtar Kelime · Desi</p>
    </div>
    """, unsafe_allow_html=True)
    guide_file = st.file_uploader(
        "Desi Kılavuzunu Yükle",
        type=["xlsx", "xls"],
        key="guide",
        label_visibility="collapsed",
    )

with col_up2:
    st.markdown("""
    <div class="upload-card">
      <h4>📦 Günlük Siparişler</h4>
      <p>Ürün Adı · Kargo Kodu</p>
    </div>
    """, unsafe_allow_html=True)
    orders_file = st.file_uploader(
        "Sipariş Dosyasını Yükle",
        type=["xlsx", "xls"],
        key="orders",
        label_visibility="collapsed",
    )

st.markdown("<hr>", unsafe_allow_html=True)

# ─────────────────────────────────────────────
# PROCESS BUTTON
# ─────────────────────────────────────────────
run_btn = st.button("⚙️ İşlemi Başlat", use_container_width=False)

if run_btn:
    # ── Guard: both files required ──
    if not guide_file or not orders_file:
        st.error("⚠️ Lütfen her iki dosyayı da yükleyin: Desi Kılavuzu ve Günlük Siparişler.")
        st.stop()

    with st.spinner("Veriler işleniyor, lütfen bekleyin..."):

        # ── Read files ──
        try:
            df_guide = pd.read_excel(guide_file)
        except Exception:
            st.error("❌ **Desi Kılavuzu** okunamadı. Geçerli bir Excel (.xlsx/.xls) dosyası yüklediğinizden emin olun.")
            st.stop()

        try:
            df_orders = pd.read_excel(orders_file)
        except Exception:
            st.error("❌ **Günlük Siparişler** dosyası okunamadı. Geçerli bir Excel (.xlsx/.xls) dosyası yüklediğinizden emin olun.")
            st.stop()

        # ── Validate columns ──
        missing_guide = validate_columns(df_guide, ["Anahtar Kelime", "Desi"], "Desi Kılavuzu")
        if missing_guide:
            st.error(
                f"❌ **Desi Kılavuzu** dosyasında şu sütunlar eksik: `{'`, `'.join(missing_guide)}`\n\n"
                "Lütfen dosyanızın **A sütununda** `Anahtar Kelime` ve **B sütununda** `Desi` başlıklarını kontrol edin."
            )
            st.stop()

        missing_orders = validate_columns(df_orders, ["Ürün Adı", "Kargo Kodu"], "Günlük Siparişler")
        if missing_orders:
            st.error(
                f"❌ **Günlük Siparişler** dosyasında şu sütunlar eksik: `{'`, `'.join(missing_orders)}`\n\n"
                "Lütfen `Ürün Adı` ve `Kargo Kodu` sütunlarının var olduğunu kontrol edin."
            )
            st.stop()

        # ── Clean & validate guide data ──
        df_guide = df_guide[["Anahtar Kelime", "Desi"]].copy()
        df_guide.dropna(subset=["Anahtar Kelime"], inplace=True)
        df_guide["Anahtar Kelime"] = df_guide["Anahtar Kelime"].astype(str).str.strip()
        df_guide["Desi"] = pd.to_numeric(df_guide["Desi"], errors="coerce")

        invalid_desi = df_guide["Desi"].isna().sum()
        if invalid_desi > 0:
            st.warning(
                f"⚠️ Desi Kılavuzunda **{invalid_desi} satırda** Desi değeri sayısal değil. "
                "Bu satırlar eşleştirmeden **çıkarıldı**. Lütfen kılavuz dosyasını kontrol edin."
            )
        df_guide.dropna(subset=["Desi"], inplace=True)

        # Duplicate keyword check
        dup_kws = df_guide[df_guide["Anahtar Kelime"].duplicated(keep=False)]["Anahtar Kelime"].unique()
        if len(dup_kws) > 0:
            st.warning(
                f"⚠️ Desi Kılavuzunda **{len(dup_kws)} tekrarlayan anahtar kelime** tespit edildi: "
                f"`{'`, `'.join(dup_kws[:5])}{'...' if len(dup_kws) > 5 else ''}`  \n"
                "Her anahtar kelime için **son satırdaki değer** kullanılmaktadır."
            )
            df_guide = df_guide.drop_duplicates(subset=["Anahtar Kelime"], keep="last")

        # Build keyword → desi map (already lowercased for matching)
        keyword_map = {
            row["Anahtar Kelime"].lower(): row["Desi"]
            for _, row in df_guide.iterrows()
        }

        # ── Clean orders data ──
        df_orders = df_orders[["Ürün Adı", "Kargo Kodu"]].copy()
        df_orders.dropna(subset=["Kargo Kodu"], inplace=True)

        # Remove fully empty product names
        df_orders["Ürün Adı"] = df_orders["Ürün Adı"].astype(str).str.strip()
        df_orders = df_orders[df_orders["Ürün Adı"] != ""]
        df_orders = df_orders[df_orders["Ürün Adı"].str.lower() != "nan"]

        # Duplicate Kargo Kodu check
        dup_codes = df_orders[df_orders["Kargo Kodu"].duplicated(keep=False)]
        if not dup_codes.empty:
            n_dup = dup_codes["Kargo Kodu"].nunique()
            st.warning(
                f"⚠️ Sipariş dosyasında **{n_dup} tekrarlayan Kargo Kodu** bulundu. "
                "Tüm tekrarlayan satırlar çıktıya dahil edilmiştir, ancak lütfen kaynaktaki veriyi doğrulayın."
            )

        # ── Perform matching ──
        df_orders["Desi"] = df_orders["Ürün Adı"].apply(
            lambda name: find_desi(name, keyword_map)
        )

        # ── Stats ──
        total       = len(df_orders)
        matched     = (df_orders["Desi"] != "KONTROL ET").sum()
        unmatched   = total - matched
        match_rate  = round((matched / total) * 100, 1) if total > 0 else 0.0

    # ─── METRICS ───
    st.markdown("### 📊 İşlem Özeti")
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("📦 Toplam Sipariş",      f"{total:,}")
    m2.metric("✅ Eşleşen",             f"{matched:,}",   delta=f"%{match_rate}")
    m3.metric("⚠️ Kontrol Edilecek",    f"{unmatched:,}", delta=f"-{unmatched}" if unmatched else "Tümü eşleşti", delta_color="inverse")
    m4.metric("🗂️ Benzersiz Anahtar",   f"{len(keyword_map):,}")

    st.markdown("<hr>", unsafe_allow_html=True)

    # ─── TABS ───
    df_matched   = df_orders[df_orders["Desi"] != "KONTROL ET"].reset_index(drop=True)
    df_unmatched = df_orders[df_orders["Desi"] == "KONTROL ET"].reset_index(drop=True)

    tab_matched, tab_unmatched, tab_all = st.tabs([
        f"✅ Eşleşenler ({matched})",
        f"⚠️ Kontrol Edilecekler ({unmatched})",
        f"📋 Tüm Sonuçlar ({total})",
    ])

    with tab_matched:
        if df_matched.empty:
            st.info("Hiçbir ürün eşleşmedi. Anahtar kelime kılavuzunuzu gözden geçirin.")
        else:
            st.caption(f"Toplam {matched} sipariş başarıyla eşleştirildi.")
            st.dataframe(df_matched, use_container_width=True, height=420)

    with tab_unmatched:
        if df_unmatched.empty:
            st.success("🎉 Tüm siparişler eşleşti! Kontrol edilecek ürün bulunmuyor.")
        else:
            st.caption(
                f"Aşağıdaki {unmatched} ürün için kılavuzda eşleşen anahtar kelime bulunamadı. "
                "Desi Kılavuzunuza eksik ürün tiplerini ekleyebilirsiniz."
            )
            st.dataframe(df_unmatched, use_container_width=True, height=420)

            # Show unique unmatched product names to help user update their guide
            with st.expander("🔍 Eşleşmeyen Benzersiz Ürün Adları (Kılavuz Güncellemesi İçin)"):
                uniq_products = df_unmatched["Ürün Adı"].drop_duplicates().sort_values()
                st.dataframe(
                    uniq_products.reset_index(drop=True).rename("Ürün Adı"),
                    use_container_width=True,
                )

    with tab_all:
        st.caption("Tüm siparişler — eşleşenler ve kontrol edilecekler birlikte.")
        # Color-code in the dataframe display
        def highlight_unmatched(row):
            if row["Desi"] == "KONTROL ET":
                return ["background-color: #78350f33; color: #f59e0b"] * len(row)
            return [""] * len(row)

        st.dataframe(
            df_orders.style.apply(highlight_unmatched, axis=1),
            use_container_width=True,
            height=460,
        )

    st.markdown("<hr>", unsafe_allow_html=True)

    # ─── EXPORT ───
    st.markdown("### 💾 Kargo Dosyasını Dışa Aktar")

    export_df = df_orders[["Kargo Kodu", "Desi"]].copy()
    excel_bytes = build_excel(export_df)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    filename  = f"kargo_desi_{timestamp}.xlsx"

    col_dl, col_info = st.columns([1, 2], gap="large")
    with col_dl:
        st.download_button(
            label="🚀 Kargo Dosyasını İndir",
            data=excel_bytes,
            file_name=filename,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    with col_info:
        st.markdown(f"""
**Dosya:** `{filename}`  
**İçerik:** `Kargo Kodu` + `Desi` — {total} satır  
**Format:** Excel (.xlsx), başlık satırı renklendirilerek, KONTROL ET satırları amber ile işaretlenerek  
""")
        if unmatched > 0:
            st.markdown(
                f"<span class='pill-warn'>⚠️ {unmatched} satır KONTROL ET olarak işaretlendi</span>",
                unsafe_allow_html=True
            )
        else:
            st.markdown(
                "<span class='pill-ok'>✅ Tüm satırlar eşleşti</span>",
                unsafe_allow_html=True
            )

    st.caption(
        f"🕐 İşlem zamanı: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')} · "
        f"Kılavuzda {len(keyword_map)} anahtar kelime kullanıldı."
    )

else:
    # ── Idle state — friendly guide ──
    st.markdown("""
<div style='
    background:#1a1d27;
    border:1px dashed #2e3350;
    border-radius:10px;
    padding:40px;
    text-align:center;
    color:#64748b;
'>
  <div style='font-size:3rem; margin-bottom:12px;'>⬆️</div>
  <div style='font-size:1rem; font-weight:600; color:#94a3b8; margin-bottom:6px;'>
    Başlamak için her iki dosyayı da yükleyin
  </div>
  <div style='font-size:0.85rem;'>
    Sol taraf → Desi Kılavuzu &nbsp;|&nbsp; Sağ taraf → Günlük Siparişler
  </div>
</div>
""", unsafe_allow_html=True)
